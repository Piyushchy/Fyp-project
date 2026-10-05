"""
Attach the speaker script in speaker_notes.py to the deck as PowerPoint
speaker notes, and write the printable version.

Run after build_deck.py:
    python add_notes.py built.pptx final.pptx script.md
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))

from speaker_notes import NOTES, as_markdown  # noqa: E402


def main() -> None:
    source, target, script = (Path(a) for a in sys.argv[1:4])

    presentation = Presentation(str(source))

    if len(presentation.slides) != len(NOTES):
        raise SystemExit(
            f"{len(presentation.slides)} slides but {len(NOTES)} notes - "
            "the two are keyed by position, so they have to match"
        )

    for number, slide in enumerate(presentation.slides, start=1):
        frame = slide.notes_slide.notes_text_frame
        frame.text = NOTES[number].strip()
        print(f"  slide {number}: {len(NOTES[number].split())} words")

    presentation.save(str(target))
    script.write_text(as_markdown(), encoding="utf-8")

    print(f"\nWrote {target}")
    print(f"Wrote {script}")


if __name__ == "__main__":
    main()
